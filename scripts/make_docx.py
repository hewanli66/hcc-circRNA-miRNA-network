# -*- coding: utf-8 -*-
"""
SR_manuscript_EN.md -> Word (.docx)
遵循 Scientific Reports 排版要求：
  - 单倍行距（line=240 auto）
  - 页脚页码（PAGE域）
  - 行号（w:lnNumType countBy=1）
  - 表格转 Word 原生表格，置于文末，表题在表前
  - 图件不嵌入（期刊要求单独提交）
  - 上标引用用真实上标格式
"""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = 'SR_manuscript_EN.md'
OUT = 'SR_manuscript_EN.docx'
doc = Document()

# ---------- 页面与默认样式 ----------
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(1)
    s.left_margin = s.right_margin = Inches(1)

st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
pf = st.paragraph_format
pf.line_spacing = 1.0# 单倍行距
pf.space_after = Pt(6)

# ---------- 行号 ----------
sectPr = doc.sections[0]._sectPr
ln = OxmlElement('w:lnNumType')
ln.set(qn('w:countBy'), '1')
ln.set(qn('w:start'), '1')
ln.set(qn('w:restart'), 'continuous')
sectPr.append(ln)

# ---------- 页脚页码 ----------
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
# fldChar begin
e1 = OxmlElement('w:fldChar')
e1.set(qn('w:fldCharType'), 'begin')
run._r.append(e1)
# instrText PAGE
e2 = OxmlElement('w:instrText')
e2.set(qn('xml:space'), 'preserve')
e2.text = ' PAGE '
run._r.append(e2)
# fldChar end
e3 = OxmlElement('w:fldChar')
e3.set(qn('w:fldCharType'), 'end')
run._r.append(e3)
run.font.name = 'Times New Roman'
run.font.size = Pt(10)


def add_para(text='', bold=False, italic=False, size=12, align=None,
             space_before=0, space_after=6, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    # 处理 **bold** 与 <sup>n</sup>
    for chunk in split_md(text):
        t, kind = chunk
        run = p.add_run(t)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold or (kind == 'b')
        run.italic = italic or (kind == 'i')
        if kind == 'sup':
            run.font.superscript = True
            run.font.size = Pt(size)
    return p


def split_md(text):
    """把含 **bold** / *italic* / <sup>n</sup> 的行拆成 (片段, 类型) """
    out = []
    pat = re.compile(r'\*\*(.+?)\*\*|\*(.+?)\*|<sup>(.+?)</sup>')
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], 'n'))
        if m.group(1) is not None:
            out.append((m.group(1), 'b'))
        elif m.group(2) is not None:
            out.append((m.group(2), 'i'))
        else:
            out.append((m.group(3), 'sup'))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], 'n'))
    return out or [(text, 'n')]


def add_table_md(block):
    """把 markdown 表格转成 Word 原生表格"""
    lines = [l for l in block.strip().split('\n') if l.strip().startswith('|')]
    rows = []
    for l in lines:
        cells = [c.strip() for c in l.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row[:ncol]):
            cell = t.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            # 表内markdown 标记清理
            v = re.sub(r'\*\*(.+?)\*\*', r'\1', val)
            v = re.sub(r'\*(.+?)\*', r'\1', v)
            run = p.add_run(v)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.bold = (i == 0)


# ---------- 解析并生成 ----------
text = open(SRC, encoding='utf-8').read()
lines = text.split('\n')

i = 0
tables = []# 收集表格，最后统一放文末
pending_table_caption = None

while i < len(lines):
    line = lines[i].rstrip()

    # 表格块
    if line.strip().startswith('|'):
        block = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i])
            i += 1
        # 数据来源表等"正文内小表"就地保留；带编号的主表与补充表移文末
        cap = pending_table_caption or ''
        keep_inline = ('Dataset' in block[0])          # Methods 里的数据来源表
        if keep_inline:
            add_table_md('\n'.join(block))
        else:
            tables.append((pending_table_caption, '\n'.join(block)))
        pending_table_caption = None
        continue

    # 分隔线
    if re.fullmatch(r'-{3,}', line.strip()):
        i += 1
        continue

    # 标题
    m = re.match(r'^(#{1,4})\s+(.*)$', line)
    if m:
        lvl, txt = len(m.group(1)), m.group(2).strip()
        # lvl==2 的 "Table 1 | xxx" / "Supplementary Table S1 |" 作为表题
        txt_clean = re.sub(r'\*\*(.+?)\*\*', r'\1', txt).strip()
        if re.match(r'^(Supplementary\s+)?Table\s', txt_clean):
            pending_table_caption = txt_clean
            add_para(txt_clean, bold=True, size=12, space_before=10, space_after=4)
            i += 1
            continue
        if lvl == 1:
            add_para(txt, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER,
                     space_before=0, space_after=12)
        elif lvl == 2:
            add_para(txt, bold=True, size=14, space_before=14, space_after=8)
        elif lvl == 3:
            add_para(txt, bold=True, size=12, space_before=10, space_after=6)
        else:
            add_para(txt, bold=True, size=12, space_before=8, space_after=4)
        i += 1
        continue

    # 引用块/待办提示
    if line.strip().startswith('>'):
        add_para(line.strip().lstrip('> '), italic=True, size=10, indent=0.3)
        i += 1
        continue

    # 列表项
    m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
    if m:
        indent = 0.3 if m.group(1) else 0.0
        marker = m.group(2)
        txt = m.group(3)
        if marker in ('-', '*'):
            add_para('•  ' + txt, size=12, indent=indent + 0.25)
        else:
            add_para(marker + '  ' + txt, size=12, indent=indent + 0.25)
        i += 1
        continue

    # 空行
    if not line.strip():
        i += 1
        continue

    add_para(line, size=12)
    i += 1

# ---------- 表格统一置于文末（表题紧接表格之前） ----------
if tables:
    doc.add_page_break()
    add_para('Tables', bold=True, size=14, space_after=4)
    add_para('', size=6, space_after=6)
    for cap, blk in tables:
        if cap:
            # 表题已在正文出现；此处再出一次以满足「表题紧接表格之前」
            add_para(cap, bold=True, size=11, space_before=10, space_after=3)
        add_table_md(blk)
        add_para('', size=6, space_after=4)

doc.save(OUT)

# ---------- 校验 ----------
d2 = Document(OUT)
ntbl = len(d2.tables)
npara = len(d2.paragraphs)
xml = d2.sections[0]._sectPr.xml
has_ln = 'lnNumType' in xml
ftr = d2.sections[0].footer.paragraphs[0]._p.xml
has_pg = 'PAGE' in ftr
print('[SAVED] %s' % OUT)
print('  段落 %d | 原生表格 %d | 行号 %s | 页码 %s'
      % (npara, ntbl, '✓' if has_ln else '✗', '✓' if has_pg else '✗'))
print('  大小 %.0f KB' % (os.path.getsize(OUT) / 1024))
